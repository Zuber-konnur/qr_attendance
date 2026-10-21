import io
import pandas as pd
from datetime import datetime
from app.database import db_session
from app.models import Attendance, AttendanceSession, User, DynamicField

class ExportService:
    @staticmethod
    def query_attendance(date_from=None, date_to=None, department="All", division="All",
                         semester="All", session_id="All", search=""):
        query = db_session.query(Attendance).join(User, Attendance.usn == User.usn).join(
            AttendanceSession, Attendance.session_id == AttendanceSession.id
        )

        if date_from:
            query = query.filter(Attendance.date >= date_from)
        if date_to:
            query = query.filter(Attendance.date <= date_to)
        if session_id and session_id != "All":
            try:
                query = query.filter(Attendance.session_id == int(session_id))
            except (ValueError, TypeError):
                pass
        if department and department != "All":
            query = query.filter(User.department == department)
        if division and division != "All":
            query = query.filter(User.division == division)
        if semester and semester != "All":
            query = query.filter(User.semester == semester)
        if search:
            search_term = f"%{search.strip()}%"
            query = query.filter((User.usn.ilike(search_term)) | (User.name.ilike(search_term)))

        return query.order_by(Attendance.date.desc(), Attendance.checkin_time.desc()).all()

    @classmethod
    def generate_dataframe(cls, records):
        dynamic_fields = db_session.query(DynamicField).filter_by(is_active=True).order_by(DynamicField.display_order).all()

        rows = []
        for r in records:
            u = r.user
            s = r.session
            custom_data = u.custom_data if (u and u.custom_data) else {}

            checkin_str = r.checkin_time.strftime("%I:%M:%S %p") if r.checkin_time else "--"
            checkout_str = r.checkout_time.strftime("%I:%M:%S %p") if r.checkout_time else "--"
            if s and s.mode == "single_checkin":
                checkout_str = "N/A (One-Time)"

            row_data = {
                "Date": r.date.strftime("%Y-%m-%d") if r.date else "",
                "Session": s.title if s else "General",
                "Mode": "One-Time Count" if (s and s.mode == "single_checkin") else "Check-in & Out",
                "USN": r.usn,
                "Student Name": u.name if u else "Unknown",
                "Department": u.department if u else "",
                "Division": u.division if u else "",
                "Semester": u.semester if u else "",
                "Check-in Time": checkin_str,
                "Check-out Time": checkout_str,
                "Duration": r.duration_formatted() if (s and s.mode != "single_checkin") else "--",
                "Status": r.status or "Present",
                "Location": r.location_name or "Authorized Area",
                "Device ID": r.device_id or "",
                "Remarks": r.remarks or ""
            }

            for df in dynamic_fields:
                row_data[df.label] = custom_data.get(df.field_key, "")

            rows.append(row_data)

        if not rows:
            cols = ["Date", "Session", "Mode", "USN", "Student Name", "Department",
                    "Division", "Semester", "Check-in Time", "Check-out Time",
                    "Duration", "Status", "Location", "Device ID", "Remarks"]
            for df in dynamic_fields:
                cols.append(df.label)
            return pd.DataFrame(columns=cols)

        return pd.DataFrame(rows)

    @classmethod
    def export_excel(cls, records, filename_prefix="Attendance_Report"):
        df = cls.generate_dataframe(records)
        output = io.BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="Attendance Data")
            worksheet = writer.sheets["Attendance Data"]

            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            for row_idx in range(2, len(df) + 2):
                is_alt = (row_idx % 2 == 0)
                for col_idx in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    if is_alt:
                        cell.fill = alt_fill

            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output.seek(0)
        return output

    @classmethod
    def export_csv(cls, records):
        df = cls.generate_dataframe(records)
        output = io.StringIO()
        df.to_csv(output, index=False)
        return output.getvalue()
